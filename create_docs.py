from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor as PRGB
import pandas as pd, json, textwrap

ROOT=Path('/mnt/data/Centrifugal_Pump_KBS')
DOCS=ROOT/'docs'; DOCS.mkdir(exist_ok=True)

# ---------- helpers ----------
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_text(cell,text,bold=False,size=9,color=None):
    cell.text=''; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None, font_size=8):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i],h,True,font_size,(255,255,255)); shade(t.rows[0].cells[i],'1F4E78')
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i],val,False,font_size)
    return t

def add_heading(doc,text,level=1):
    p=doc.add_heading(text,level=level)
    return p

def add_bullets(doc, items, level=0):
    for it in items:
        p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(it)

def add_code(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25)
    for line in text.strip().splitlines():
        r=p.add_run(line+'\n'); r.font.name='Consolas'; r.font.size=Pt(8.5)

# ---------- Report ----------
doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(0.65); sec.bottom_margin=Inches(0.65); sec.left_margin=Inches(0.7); sec.right_margin=Inches(0.7)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
for s in ['Title','Heading 1','Heading 2','Heading 3']:
    styles[s].font.name='Aptos Display'

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('DEPARTMENT OF COMPUTER SCIENCE & INFORMATION TECHNOLOGY'); r.bold=True; r.font.size=Pt(13)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('BACHELORS OF SCIENCE IN COMPUTER SCIENCE'); r.bold=True; r.font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(18)
r=p.add_run('Knowledge Based System (CT-359)'); r.font.size=Pt(12)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Knowledge-Based Expert System for Centrifugal Pump\nFault Diagnosis and Maintenance Recommendation'); r.bold=True; r.font.size=Pt(20); r.font.color.rgb=RGBColor(31,78,120)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('\nProject Report\n\nGroup Members: ______________________________\nInstructor: _________________________________\nSubmission Date: 18 October 2026').font.size=Pt(11)
doc.add_page_break()

add_heading(doc,'Abstract',1)
doc.add_paragraph('This project develops an explainable Knowledge-Based System (KBS) for diagnosing faults in a centrifugal pump and recommending maintenance actions. The prototype represents heterogeneous equipment knowledge through weighted production rules linking components, symptoms, faults, likely causes, safety priorities, repair cost, downtime, and maintenance actions. Numeric operating observations are transformed into symbolic evidence and processed using forward-chaining style rule evaluation plus weighted evidence ranking. The system returns multiple candidate faults with confidence-style evidence scores, an explanation trace, and prioritized maintenance recommendations. The project includes 14 fault classes, 22 symptoms, 16 maintenance actions, 14 compound rules, 280 synthetic diagnostic cases, and a 255-record synthetic sensor log. On the internally generated diagnostic cases, the prototype achieved 85.7% top-1 and 98.2% top-3 agreement with the generating label; these figures demonstrate internal consistency only and are not a claim of real-world diagnostic accuracy.')

add_heading(doc,'1. Problem Definition and Objectives',1)
doc.add_paragraph('Industrial pump knowledge is typically distributed across maintenance manuals, operator observations, historical records, expert experience, and sensor readings. The proposed KBS brings these sources into a structured knowledge model so that maintenance personnel can enter observed evidence and receive a transparent diagnostic recommendation.')
add_bullets(doc,[
'Represent equipment, components, symptoms, faults, causes, evidence, and maintenance actions in a machine-readable form.',
'Analyze stakeholder needs, incomplete requirements, operational constraints, and safety priorities.',
'Compare alternative knowledge-representation and inference approaches and justify the selected method.',
'Implement a working prototype capable of handling incomplete or overlapping evidence.',
'Provide explainable ranked diagnoses and maintenance recommendations.'
])

add_heading(doc,'2. CLO and Complex-Problem Mapping',1)
rows=[
('CLO 1 – Knowledge Representation','Components, symptoms, faults, maintenance actions, weighted relationships, and production rules are stored in CSV/JSON knowledge bases.'),
('CLO 2 – Problem Analysis & Solution Design','Stakeholders, functional/non-functional requirements, incomplete information, safety, downtime, cost, and production continuity are analyzed and reflected in the architecture.'),
('CLO 3 – Method & Algorithm Assessment','Rule-based, ontology/knowledge-graph, case-based, and data-driven approaches are compared. Weighted forward reasoning is selected for transparency and manageable data requirements.'),
('CPA-2','There is no single obvious diagnosis when multiple faults share symptoms; the system ranks competing hypotheses.'),
('CPA-3','The solution combines KBS concepts, pump domain knowledge, Python implementation, structured data, and reasoning.'),
('CPA-9','Unclear sensor thresholds, missing observations, inconsistent evidence, and stakeholder trade-offs are explicitly identified.')]
add_table(doc,['Outcome / Attribute','How addressed'],rows,font_size=8.5)

add_heading(doc,'3. Selected Equipment: Centrifugal Pump',1)
doc.add_paragraph('A centrifugal pump was selected because it has well-defined hydraulic, mechanical, electrical, and sealing subsystems, and many faults exhibit overlapping but explainable symptom patterns. This makes it suitable for a knowledge-based diagnostic prototype.')
add_table(doc,['Component','Category','Role'],[(x[1],x[2],x[3]) for x in [tuple(r) for r in pd.read_csv(ROOT/'data/components.csv').itertuples(index=False,name=None)]],font_size=8)

add_heading(doc,'4. Stakeholders and Requirements',1)
add_heading(doc,'4.1 Stakeholders',2)
add_table(doc,['Stakeholder','Primary need'],[
('Maintenance technician','Likely fault, cause, inspection steps, repair priority, explanation.'),
('Operator','Simple symptom entry and immediate warning when a dangerous pattern appears.'),
('Maintenance engineer','Traceable rules, editable knowledge base, alternative diagnoses.'),
('Production manager','Reduced downtime and visibility of maintenance impact.'),
('Safety personnel','Conservative warnings and safe isolation recommendation for critical conditions.'),
('Management','Lower unplanned maintenance cost and consistent decision support.')],font_size=8.5)
add_heading(doc,'4.2 Functional Requirements',2)
add_bullets(doc,['Accept numeric sensor observations and manual symptoms.','Convert numeric readings into symbolic evidence.','Match evidence against stored fault knowledge.','Rank multiple possible faults instead of forcing a single answer.','Show confidence/evidence score, safety priority, cost, and downtime.','Recommend maintenance actions.','Show which evidence and rules caused each conclusion.'])
add_heading(doc,'4.3 Non-Functional Requirements',2)
add_bullets(doc,['Explainability and traceability.','Fast interactive response.','Knowledge base should be editable without rewriting the entire program.','Clear separation between knowledge base and inference engine.','Educational safety warning: prototype thresholds are not manufacturer limits.'])
add_heading(doc,'4.4 Incomplete / Ambiguous Requirements',2)
add_bullets(doc,['The assignment does not specify pump model, duty point, fluid properties, motor rating, sensor type, or OEM alarm limits.','Real diagnostic thresholds vary by machine and service; therefore the prototype uses clearly labeled academic demo thresholds.','Conflicting evidence can occur because faults share symptoms such as vibration, low flow, and heating.','Not every symptom is always observable; therefore the inference engine uses partial evidence rather than requiring a complete rule match.'])

add_heading(doc,'5. Knowledge Representation Design',1)
doc.add_paragraph('The project uses a hybrid rule-based representation. Atomic relationships are stored as weighted fault–symptom links. Higher-order patterns are represented as production rules whose conditions must all be observed before the rule fires. Metadata for safety, repair cost, downtime, and maintenance actions is stored separately so recommendations remain modular.')
add_table(doc,['Entity','Examples','Storage'],[
('Equipment / components','Impeller, bearing, seal, coupling, suction line, motor','components.csv'),
('Symptoms / evidence','High vibration, low flow, seal leakage, crackling noise','symptoms.csv'),
('Faults','Cavitation, bearing failure, misalignment, dry running','faults.csv'),
('Weighted relationships','Fault F02 ↔ suction_pressure_low weight 5','symptom_fault_weights.csv'),
('Compound rules','IF crackling_noise AND suction_pressure_low AND pressure_fluctuation THEN cavitation','rules.json'),
('Maintenance actions','Inspect bearing, restore prime, clean strainer','maintenance_actions.csv')],font_size=8)

add_heading(doc,'6. Fault Knowledge Base',1)
f=pd.read_csv(ROOT/'data/faults.csv')
add_table(doc,['ID','Fault','Priority','Cost','Downtime'],f[['fault_id','fault_name','safety_priority','repair_cost','downtime']].values.tolist(),font_size=7.6)

add_heading(doc,'7. Example Production Rules',1)
add_code(doc,'''
R01: IF vibration_high AND bearing_temp_high AND grinding_noise
     THEN Bearing Wear / Failure

R02: IF crackling_noise AND suction_pressure_low AND pressure_fluctuation
     THEN Cavitation

R07: IF no_flow AND prime_failure AND rapid_temp_rise
     THEN Dry Running / Loss of Prime

R12: IF air_bubbles AND flow_fluctuation AND pressure_fluctuation
     THEN Air Ingress on Suction Side
''')

add_heading(doc,'8. Inference Mechanism',1)
doc.add_paragraph('The inference process is forward and evidence-driven. The user provides facts. Numeric values are first converted into symbolic evidence. Each candidate fault receives points for matched evidence. If every condition of a compound rule is satisfied, its bonus is added. Candidate faults are then ranked by normalized score, followed by safety priority as a tie-breaker.')
doc.add_paragraph('For a fault f, the prototype computes:')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Score(f) = Σ matched symptom weights + Σ fired rule bonuses'); r.bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Evidence Score (%) = 100 × Score(f) / Maximum configured score(f)'); r.bold=True

doc.add_paragraph('This percentage is an explainable evidence score for the prototype, not a calibrated probability. The design is intentional: it allows incomplete evidence and competing diagnoses without hiding the reasoning in a black-box model.')

add_heading(doc,'9. Why This Method Was Selected',1)
add_table(doc,['Method','Strength','Limitation','Decision'],[
('Rule-based / forward reasoning','Highly explainable; works with expert knowledge; little training data required','Rule maintenance can grow with system size','Selected'),
('Backward chaining','Good when testing one known hypothesis','Less natural for broad symptom-first screening','Discussed, not primary'),
('Ontology / knowledge graph','Excellent semantic relationships and extensibility','More tooling and modeling overhead for this prototype','Future extension'),
('Case-based reasoning','Useful with historical maintenance cases','Requires a meaningful library of validated cases','Future extension'),
('Machine learning','Can model complex patterns from large datasets','Needs reliable labeled data; explanation may be weaker','Not selected for core inference')],font_size=7.8)

add_heading(doc,'10. System Architecture',1)
add_code(doc,'''
Operator / Sensor Inputs
         |
         v
Numeric-to-Symbolic Evidence Layer
         |
         v
+---------------- Knowledge Base ----------------+
| Components | Symptoms | Faults | Weights | Rules|
| Safety | Cost | Downtime | Maintenance Actions  |
+--------------------------------------------------+
         |
         v
Inference Engine
(weighted matching + forward rule firing)
         |
         v
Ranked Faults -> Explanation Trace -> Maintenance Recommendation
''')

add_heading(doc,'11. Dataset and Logging Design',1)
add_table(doc,['File','Size / contents','Purpose'],[
('diagnostic_cases.csv','280 labeled synthetic cases × 22 symptom fields','Tests incomplete/overlapping symbolic evidence.'),
('operational_sensor_log.csv','255 five-minute synthetic records','Demonstrates time-stamped vibration, temperature, pressure, flow, current and labeled fault episodes.'),
('evaluation_results.csv','280 predictions','Stores top prediction and top-3 inclusion result for each diagnostic case.'),
('components.csv','10 components','Equipment structure.'),
('faults.csv','14 faults','Fault metadata and priorities.'),
('symptoms.csv','22 symptoms','Controlled vocabulary for evidence.'),
('maintenance_actions.csv','16 actions','Repair/inspection guidance.'),
('rules.json','14 compound rules','Explicit explainable IF–THEN patterns.')],font_size=8)
doc.add_paragraph('Important: the diagnostic cases and sensor logs are synthetic and reproducible. They were created for academic prototyping, software testing, and demonstration. They are not real maintenance records and must not be presented as field measurements.')

add_heading(doc,'12. Prototype Thresholds',1)
doc.add_paragraph('The software contains demo thresholds to convert numeric observations into symbolic evidence, for example vibration ≥ 6.0 mm/s RMS → vibration_high, bearing temperature ≥ 80°C → bearing_temp_high, and flow < 75% of a healthy baseline → flow_low. These values are implementation defaults for the academic demonstration only. A real deployment would load model-specific operating envelopes and alarm limits from the manufacturer or plant engineering database.')

add_heading(doc,'13. Evaluation',1)
doc.add_paragraph('The inference engine was evaluated on the 280 internally generated symbolic diagnostic cases. The result was 85.7% top-1 agreement and 98.2% top-3 inclusion of the generating fault label. Because the cases were generated from the same knowledge model, this evaluation checks implementation consistency rather than independent field validity. Real accuracy must be validated on expert-reviewed maintenance cases from the target plant.')
add_table(doc,['Metric','Result','Interpretation'],[
('Top-1 agreement','85.7%','The highest-ranked diagnosis matched the synthetic generating label in 240/280 cases.'),
('Top-3 inclusion','98.2%','The generating label appeared among the first three candidates in 275/280 cases.'),
('Automated unit tests','5/5 passed','Representative bearing, cavitation, dry-running, air-ingress, and sensor-conversion cases passed.')],font_size=8.5)

add_heading(doc,'14. Demonstration Scenarios',1)
add_table(doc,['Scenario','Evidence','Expected top diagnosis'],[
('Bearing distress','High vibration + hot bearing + grinding/rumbling noise','Bearing Wear / Failure'),
('Cavitation','Crackling noise + low suction pressure + pressure fluctuation + low flow','Cavitation'),
('Dry running','Near-zero flow + loss of prime + rapid temperature rise','Dry Running / Loss of Prime'),
('Air ingress','Air bubbles + unstable pressure/flow','Air Ingress on Suction Side'),
('Discharge restriction','High discharge pressure + low flow + rapid heating','Discharge Restriction / Closed Valve')],font_size=8.2)

add_heading(doc,'15. Safety, Cost, Downtime and Explainability',1)
doc.add_paragraph('Each fault is associated with safety priority, qualitative repair cost, and qualitative downtime. The interface highlights critical conditions, while the explanation trace shows matched evidence and rule identifiers. This supports accountable decision-making: the user can see not only what the system recommends, but why the recommendation appeared.')

add_heading(doc,'16. Limitations and Future Work',1)
add_bullets(doc,[
'Knowledge is manually engineered and limited to the selected centrifugal-pump fault set.',
'The evidence score is not a statistical probability and is not calibrated on field data.',
'Synthetic logs do not capture all real process disturbances, simultaneous faults, sensor bias, or transient behavior.',
'A real deployment should integrate OEM-specific thresholds, historian/SCADA data, maintenance work orders, and expert validation.',
'Future versions can add an ontology/knowledge graph, temporal rules, case-based reasoning, probabilistic confidence, and live sensor integration.'
])

add_heading(doc,'17. Conclusion',1)
doc.add_paragraph('The project demonstrates how a KBS can transform fragmented industrial maintenance knowledge into a structured, explainable fault-diagnosis tool. The weighted production-rule approach is appropriate for the assignment because it supports knowledge representation, problem analysis, method comparison, incomplete evidence, alternative diagnoses, and transparent maintenance recommendations while remaining understandable enough to defend during a viva.')

add_heading(doc,'References',1)
refs=[
'[1] KSB, “Maintenance – Centrifugal Pump Lexicon.” https://www.ksb.com/en-global/centrifugal-pump-lexicon/article/maintenance-1116120',
'[2] KSB, “Dry running – Centrifugal Pump Lexicon.” https://www.ksb.com/en-global/centrifugal-pump-lexicon/article/dry-running-1116984',
'[3] KSB, “Operating behaviour – Centrifugal Pump Lexicon.” https://www.ksb.com/en-global/centrifugal-pump-lexicon/article/operating-behaviour-1118222',
'[4] KSB, “Pipeline pump – Centrifugal Pump Lexicon.” https://www.ksb.com/en-global/centrifugal-pump-lexicon/article/pipeline-pump-1116452',
'[5] Project source files and synthetic datasets supplied with this submission.'
]
for ref in refs: doc.add_paragraph(ref)

report_path=DOCS/'CT359_Centrifugal_Pump_KBS_Report.docx'; doc.save(report_path)

# ---------- Viva guide ----------
viva='''# CT-359 Centrifugal Pump KBS — Viva Guide\n\n## 30-second project explanation\nOur project is an explainable knowledge-based expert system for centrifugal pump fault diagnosis. The knowledge base stores components, symptoms, faults, symptom weights, compound rules, safety levels, cost/downtime, and maintenance actions. The inference engine converts sensor values into symbolic facts, applies weighted forward reasoning, fires matching IF–THEN rules, ranks possible faults, and explains which evidence produced the result.\n\n## Most important distinction\n**Knowledge base = what the system knows. Inference engine = how the system reasons with that knowledge.**\n\n## Why centrifugal pump?\nIt combines hydraulic, mechanical, sealing, and electrical faults, and different faults share symptoms. That makes it a good complex-problem example for knowledge representation and explainable reasoning.\n\n## Why rules instead of machine learning?\nRules are transparent, easy to validate with domain experts, and do not require a large labeled dataset. Machine learning can be future work after collecting validated field history.\n\n## Why forward chaining?\nWe start from known facts such as vibration, temperature, pressure, flow, noise, and leakage, and infer possible faults. That is a data-driven direction, so forward reasoning is natural.\n\n## How do you handle incomplete information?\nA full rule match is not mandatory. Each observed symptom contributes a weight to candidate faults. Compound rule bonuses are added only when all their conditions are available. Therefore partial evidence still produces a ranked result.\n\n## Is the confidence a probability?\nNo. It is an explainable normalized evidence score, not a statistically calibrated probability. This is an important limitation.\n\n## Are the logs real?\nNo. They are synthetic educational logs generated for prototype testing. We state this clearly. Real deployment would require OEM/plant data and expert validation.\n\n## How is conflicting evidence handled?\nSeveral faults can receive evidence at the same time. The system ranks them instead of forcing one answer and displays the top alternatives with their matching evidence.\n\n## What makes the system explainable?\nFor every result it shows matched symptoms, fired rule IDs, rule explanations, safety priority, and recommended maintenance actions.\n\n## Main dataset numbers\n- 10 components\n- 14 fault classes\n- 22 symbolic symptoms\n- 16 maintenance actions\n- 14 compound rules\n- 280 synthetic diagnostic cases\n- 255 time-stamped sensor-log records\n\n## Internal test result\nTop-1 agreement: 85.7%. Top-3 inclusion: 98.2%. This checks internal consistency only because the cases are synthetic and based on the same knowledge model.\n\n## Best demo cases\n1. Bearing failure: vibration 8.5, bearing temp 90, grinding noise ON.\n2. Cavitation: vibration 7, flow 70, discharge pressure 4, suction pressure 0.5, fluctuation 15, crackling noise ON.\n3. Dry running: flow 3, discharge pressure 0.5, suction pressure 0.2, seal temp 100, pump primed OFF, rapid temperature rise ON.\n\n## Questions you may be asked\n**What is a production rule?** An IF conditions THEN conclusion representation of expert knowledge.\n\n**What is knowledge acquisition?** Collecting and structuring knowledge from manuals, experts, maintenance history, and operational data.\n\n**Why not ontology?** It is strong for semantic relationships, but production rules are simpler for this prototype and directly support explainable diagnosis. An ontology can be added later.\n\n**What is a limitation of rule-based systems?** Rule bases can become difficult to maintain as the domain expands, and poorly designed rules can conflict.\n\n**What would you add in industry?** Real SCADA/historian inputs, OEM alarm limits, work-order history, expert validation, temporal reasoning, and possibly probabilistic or machine-learning support.\n'''
(DOCS/'Viva_Guide.md').write_text(viva,encoding='utf-8')

# ---------- PPTX ----------
prs=Presentation(); prs.slide_width=PInches(13.333); prs.slide_height=PInches(7.5)
TITLE=PRGB(31,78,120); DARK=PRGB(30,30,30); LIGHT=PRGB(242,246,250); ACC=PRGB(47,117,181)

def add_title(slide,title,subtitle=None):
    box=slide.shapes.add_textbox(PInches(.65),PInches(.35),PInches(12),PInches(.7))
    p=box.text_frame.paragraphs[0]; p.text=title; p.font.size=PPt(26); p.font.bold=True; p.font.color.rgb=TITLE
    if subtitle:
        s=slide.shapes.add_textbox(PInches(.67),PInches(1.0),PInches(11.8),PInches(.45)); sp=s.text_frame.paragraphs[0]; sp.text=subtitle; sp.font.size=PPt(11); sp.font.color.rgb=PRGB(90,90,90)

def add_footer(slide,n):
    t=slide.shapes.add_textbox(PInches(11.8),PInches(7.1),PInches(.9),PInches(.25)); p=t.text_frame.paragraphs[0]; p.text=str(n); p.font.size=PPt(9); p.font.color.rgb=PRGB(120,120,120); p.alignment=PP_ALIGN.RIGHT

def bullets(slide,items,x=.8,y=1.5,w=11.7,h=5.2,size=20):
    tb=slide.shapes.add_textbox(PInches(x),PInches(y),PInches(w),PInches(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.clear()
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=it; p.font.size=PPt(size); p.font.color.rgb=DARK; p.space_after=PPt(8); p.level=0

def title_slide(title,sub):
    s=prs.slides.add_slide(prs.slide_layouts[6]);
    shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height); shape.fill.solid(); shape.fill.fore_color.rgb=LIGHT; shape.line.fill.background()
    box=s.shapes.add_textbox(PInches(1),PInches(1.8),PInches(11.3),PInches(1.8)); p=box.text_frame.paragraphs[0]; p.text=title; p.font.size=PPt(34); p.font.bold=True; p.font.color.rgb=TITLE; p.alignment=PP_ALIGN.CENTER
    sb=s.shapes.add_textbox(PInches(1.5),PInches(4.2),PInches(10.3),PInches(1.1)); p=sb.text_frame.paragraphs[0]; p.text=sub; p.font.size=PPt(18); p.font.color.rgb=DARK; p.alignment=PP_ALIGN.CENTER
    return s

slides=[]
slides.append(title_slide('Centrifugal Pump Fault Diagnosis KBS','CT-359 Knowledge Based System | Group Members: __________'))
content=[
('Problem & Goal',['Industrial maintenance knowledge is scattered across manuals, observations, history, and sensors.','Goal: convert this knowledge into an explainable decision-support system.','Input: symptoms + operating evidence. Output: ranked faults + maintenance actions + explanation.']),
('CLO Mapping',['CLO 1: formal knowledge representation using symptoms, faults, components, rules, and actions.','CLO 2: stakeholder and requirement analysis, safety/cost/downtime constraints.','CLO 3: compare reasoning methods and justify weighted forward rule reasoning.']),
('Why Centrifugal Pump?',['Contains hydraulic, mechanical, sealing, structural, and electrical subsystems.','Faults share symptoms, so diagnosis is not a simple one-to-one lookup.','Good domain for explainable reasoning and competing fault hypotheses.']),
('Knowledge Base',['10 components','14 fault classes','22 symbolic symptoms','16 maintenance actions','14 compound IF–THEN rules','Weighted fault–symptom relationships']),
('Key Fault Classes',['Bearing wear/failure','Cavitation','Shaft misalignment','Impeller blockage/fouling','Mechanical seal failure','Suction blockage','Dry running / loss of prime','Motor overload','Air ingress','Discharge restriction']),
('Knowledge Representation',['CSV files store structured entities and relationships.','JSON stores compound production rules.','Weights represent relative diagnostic importance of symptoms.','Fault metadata stores safety priority, repair cost, and expected downtime.']),
('Example Rules',['R01: vibration_high AND bearing_temp_high AND grinding_noise → Bearing failure','R02: crackling_noise AND suction_pressure_low AND pressure_fluctuation → Cavitation','R07: no_flow AND prime_failure AND rapid_temp_rise → Dry running']),
('Inference Mechanism',['1. Convert numeric sensor readings into symbolic evidence.','2. Add points for matched weighted symptoms.','3. Fire compound rules when all conditions are present.','4. Normalize into an evidence score.','5. Rank candidate faults and attach maintenance actions.']),
('Handling Incomplete / Conflicting Evidence',['Full rule match is not required.','Partial symptoms still contribute weighted evidence.','Multiple faults may be returned simultaneously.','System exposes alternatives instead of hiding ambiguity.']),
('System Architecture',['Inputs → evidence conversion → knowledge base → inference engine → ranked diagnosis.','Knowledge base is separated from reasoning code.','Explanation layer returns matched evidence and rule IDs.']),
('Dataset & Logs',['280 synthetic labeled diagnostic cases.','255 five-minute synthetic operating records.','Sensors include vibration, bearing/motor/seal temperature, flow, pressure, current, fluctuation.','Synthetic data is clearly labeled; it is not real plant history.']),
('Evaluation',['Top-1 internal agreement: 85.7%.','Top-3 inclusion: 98.2%.','5/5 automated unit tests passed.','Important: this demonstrates implementation consistency, not independent field accuracy.']),
('Demo: Bearing Failure',['Enter vibration ≈ 8.5 mm/s RMS.','Bearing temperature ≈ 90°C.','Turn grinding/rumbling noise ON.','Expected top result: Bearing Wear / Failure with rule R01.']),
('Demo: Cavitation',['Use low suction pressure, low flow, pressure fluctuation, crackling noise.','Expected top result: Cavitation.','System recommends checking suction conditions and NPSH-related causes.']),
('Safety, Cost & Explainability',['Critical patterns receive stronger warnings.','Recommendations include qualitative cost and downtime.','Every diagnosis shows evidence and fired rules.','Prototype is educational and not a substitute for OEM/plant safety procedures.']),
('Limitations & Future Work',['Use OEM/plant-specific thresholds.','Validate on expert-reviewed real maintenance cases.','Add ontology/knowledge graph and temporal reasoning.','Connect SCADA/historian data and work-order records.','Add probabilistic confidence or ML only after data validation.']),
('Conclusion',['A complete explainable KBS prototype was implemented.','It directly addresses knowledge representation, problem analysis, and method assessment.','The design is sufficiently complex for the project while remaining defendable in viva.'])]
for idx,(title,items) in enumerate(content,2):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_title(s,title); bullets(s,items); add_footer(s,idx)

ppt_path=DOCS/'CT359_Centrifugal_Pump_KBS_Presentation.pptx'; prs.save(ppt_path)
print(report_path); print(ppt_path)
